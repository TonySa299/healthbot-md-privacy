"""
Write each new daily entry back into the Excel workbook, so the spreadsheet
stays in sync with what is entered in the app.

Design choices for safety (this is the owner's master file):
  * A timestamped backup is taken before every write (last 20 kept).
  * The workbook is opened with formulas preserved (data_only=False).
  * New rows are appended after the last data row of the "Daily <Station>" and
    "ODOMETERS <Station>" sheets; contiguous Excel tables are grown by one row
    so their formulas keep covering the data. Values (not formulas) are written,
    matching exactly what the app computed.
  * Any failure is caught by the caller so a sync problem never blocks saving
    to the app's own database.
"""
import os
import glob
import shutil
import datetime

import openpyxl
from openpyxl.utils import get_column_letter, range_boundaries

from db import LIVE_WORKBOOK, BACKUP_DIR

# Header name (lower-cased) -> the record key produced by the app.
DAILY_COLMAP = {
    "date": "day", "price change": "price_note",
    "benzine 1 l": "b1_liters", "benzine 1 p": "b1_price", "total benzine 1": "b1_total",
    "benzine 2 l": "b2_liters", "benzine 2 p": "b2_price", "total benzine 2": "b2_total",
    "total benzine l (1+2)": "benzine_liters", "total benzine $": "benzine_total",
    "mezout l": "mezout_liters", "mezout p": "mezout_price", "total mezout": "mezout_total",
    "oil": "oil", "wash": "wash", "gaz": "gaz", "fragrence": "fragrance",
    "kiosk": "kiosk", "payments": "payments", "coupon": "coupon_in", "total in": "total_in",
    "debts": "debts", "internal sale": "internal_sale", "coupons": "coupons",
    "difference": "difference",
    "purchase oil": "purchase_oil", "purchase kiosk": "purchase_kiosk",
    "purchase gaz": "purchase_gaz", "other purchases": "other_purchases",
    "total purchases": "total_purchases",
    "pmnt - rent": "rent", "salaries": "salaries", "r&m": "repairs",
    "utilities": "utilities", "client satisfaction": "client_satisfaction",
    "tips": "tips", "new assets": "new_assets", "other": "other_expense",
    "total expense": "total_expense", "profit sharing": "profit_sharing",
    "total out": "total_out", "daily total": "daily_total", "notes": "notes",
}


def _norm(v):
    return str(v).strip().lower() if v is not None else ""


def _backup():
    if not os.path.exists(LIVE_WORKBOOK):
        return
    os.makedirs(BACKUP_DIR, exist_ok=True)
    stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    shutil.copy2(LIVE_WORKBOOK, os.path.join(BACKUP_DIR, f"General_Cashflow_{stamp}.xlsx"))
    backups = sorted(glob.glob(os.path.join(BACKUP_DIR, "General_Cashflow_*.xlsx")))
    for old in backups[:-20]:                      # keep the 20 most recent
        try:
            os.remove(old)
        except OSError:
            pass


def _header_row(ws, label="date"):
    for r in range(1, 12):
        if _norm(ws.cell(r, 1).value) == label:
            return r
    return None


def _last_data_row(ws, date_col, start):
    last = start
    for r in range(start + 1, ws.max_row + 2):
        if isinstance(ws.cell(r, date_col).value, datetime.datetime):
            last = r
    return last


def _grow_tables(ws, new_row):
    """Extend any table that ends exactly one row above the new row."""
    for tbl in ws.tables.values():
        min_c, min_r, max_c, max_r = range_boundaries(tbl.ref)
        if max_r == new_row - 1:
            tbl.ref = (f"{get_column_letter(min_c)}{min_r}:"
                       f"{get_column_letter(max_c)}{new_row}")


def _write(ws, row, col, value, copy_fmt_from=None):
    cell = ws.cell(row, col)
    cell.value = value
    if copy_fmt_from is not None:
        cell.number_format = ws.cell(copy_fmt_from, col).number_format


def append_day(station, record, odo, litres):
    """Append one day to the Daily and ODOMETERS sheets. Returns the file path.

    record: the dict of computed daily values (keys match DAILY_COLMAP values)
    odo:    dict with b1_a,b1_b,b2_a,b2_b,mez_a,mez_b
    litres: dict with b1,b2,mez (litres sold)
    """
    _backup()
    wb = openpyxl.load_workbook(LIVE_WORKBOOK, data_only=False)
    day = datetime.datetime.fromisoformat(record["day"])

    # ---- Daily <Station> -------------------------------------------------
    ws = wb[f"Daily {station}"]
    hr = _header_row(ws)
    if hr:
        colmap = {}
        for c in range(1, ws.max_column + 1):
            key = DAILY_COLMAP.get(_norm(ws.cell(hr, c).value))
            if key and key not in colmap:
                colmap[key] = c
        date_col = colmap.get("day", 1)
        last = _last_data_row(ws, date_col, hr)
        new = last + 1
        for key, c in colmap.items():
            if key == "day":
                _write(ws, new, c, day, copy_fmt_from=last)
            elif key in ("price_note", "notes"):
                val = record.get(key)
                if val:
                    _write(ws, new, c, str(val))
            else:
                _write(ws, new, c, round(float(record.get(key, 0) or 0), 6), copy_fmt_from=last)
        _grow_tables(ws, new)

    # ---- ODOMETERS <Station> --------------------------------------------
    wo = wb[f"ODOMETERS {station}"]
    hr = _header_row(wo)
    if hr:
        last = _last_data_row(wo, 1, hr)
        new = last + 1
        _write(wo, new, 1, day, copy_fmt_from=last)                 # A date
        if record.get("price_note"):
            _write(wo, new, 2, str(record["price_note"]))           # B note
        _write(wo, new, 3, odo["b1_a"], last); _write(wo, new, 4, odo["b1_b"], last)  # C,D
        _write(wo, new, 5, odo["b2_a"], last); _write(wo, new, 6, odo["b2_b"], last)  # E,F
        _write(wo, new, 7, odo["mez_a"], last); _write(wo, new, 8, odo["mez_b"], last)  # G,H
        _write(wo, new, 10, litres["b1"], last)                     # J Benzine 1
        _write(wo, new, 11, litres["b2"], last)                     # K Benzine 2
        _write(wo, new, 12, litres["mez"], last)                    # L Mezout
        _write(wo, new, 13, litres["b1"] + litres["b2"] + litres["mez"], last)  # M total
        _grow_tables(wo, new)

    wb.save(LIVE_WORKBOOK)
    return LIVE_WORKBOOK
